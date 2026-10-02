"""Safe local file persistence and supported document text extraction."""

from __future__ import annotations

import json
import ipaddress
import re
import socket
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_URL_HTML_BYTES = 5 * 1024 * 1024
URL_FETCH_TIMEOUT_SECONDS = 15
MAX_URL_REDIRECTS = 5
SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx"}


class ImportError(ValueError):
    """Raised when an uploaded document cannot be safely imported."""


def _validate_public_url(value: str) -> str:
    parsed = urlsplit(value.strip())
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
        raise ImportError("Use a valid public HTTP or HTTPS vacancy URL")
    if parsed.username or parsed.password or parsed.fragment:
        raise ImportError("Vacancy URL must not contain credentials or a fragment")
    host = parsed.hostname.rstrip(".").lower()
    try:
        port = parsed.port or (443 if parsed.scheme.lower() == "https" else 80)
        addresses = {item[4][0] for item in socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)}
    except (OSError, ValueError) as exc:
        raise ImportError("The vacancy URL host could not be resolved") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast or ip.is_reserved or ip.is_unspecified:
            raise ImportError("The vacancy URL points to a blocked internal address")
    return parsed.geturl()


class _SafeRedirectHandler(HTTPRedirectHandler):
    def __init__(self) -> None:
        super().__init__()
        self.redirects = 0

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        self.redirects += 1
        if self.redirects > MAX_URL_REDIRECTS:
            raise ImportError("The vacancy page redirected too many times")
        safe = _validate_public_url(urljoin(req.full_url, newurl))
        return super().redirect_request(req, fp, code, msg, headers, safe)


def fetch_vacancy_url(value: str) -> tuple[str, str, bytes, VacancyExtraction]:
    url = _validate_public_url(value)
    handler = _SafeRedirectHandler()
    opener = build_opener(handler)
    request = Request(url, headers={"Accept": "text/html,application/xhtml+xml", "User-Agent": "CareerFlow vacancy importer"})
    try:
        with opener.open(request, timeout=URL_FETCH_TIMEOUT_SECONDS) as response:
            final_url = _validate_public_url(response.geturl())
            content_type = response.headers.get_content_type().lower()
            if content_type not in {"text/html", "application/xhtml+xml"}:
                raise ImportError("The vacancy page did not return HTML")
            content = bytearray()
            while True:
                chunk = response.read(min(64 * 1024, MAX_URL_HTML_BYTES + 1 - len(content)))
                if not chunk:
                    break
                content.extend(chunk)
                if len(content) > MAX_URL_HTML_BYTES:
                    raise ImportError("The vacancy page is larger than the 5 MB limit")
    except ImportError:
        raise
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise ImportError("The vacancy page could not be downloaded") from exc
    extraction = extract_vacancy_html(bytes(content))
    if len(extraction.candidate_text) < 40:
        raise ImportError("The vacancy page does not contain enough readable text")
    return url, final_url, bytes(content), extraction


@dataclass(frozen=True, slots=True)
class VacancyExtraction:
    raw_text: str
    candidate_text: str
    source_type: str
    method: str


class _VisibleTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self.hidden_depth += 1
        elif tag.lower() in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self.hidden_depth = max(0, self.hidden_depth - 1)
        elif tag.lower() in {"p", "div", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.hidden_depth == 0:
            self.parts.append(data)

    def text(self) -> str:
        return normalize_text(" ".join(self.parts))


class _VacancyHTMLParser(HTMLParser):
    """Collect generic visible text and site-independent structured candidates."""

    _BLOCK_TAGS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr", "section"}
    _VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
    _HIDDEN_TAGS = {"script", "style", "noscript", "svg"}
    _SUPPRESSED_TAGS = {"header", "footer", "nav", "aside", "form", "button"}
    _SUPPRESSED_HINT = re.compile(
        r"(?i)(cookie|consent|breadcrumb|recommend|related.jobs|social|advert|login|search|menu|toolbar)"
    )
    _CONTAINER_HINT = re.compile(
        r"(?i)(job.?description|vacancy.?description|description|job.?details|posting.?content|vacancy.?content)"
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.visible: list[str] = []
        self.candidates: dict[str, list[str]] = {
            "microdata": [],
            "rdfa": [],
            "main": [],
            "article": [],
            "container": [],
        }
        self._stack: list[dict[str, object]] = []
        self._active = {key: 0 for key in self.candidates}
        self._hidden_depth = 0
        self._suppressed_depth = 0
        self._micro_job_depth = 0
        self._rdfa_job_depth = 0
        self._json_script: list[str] | None = None
        self._json_script_kind = ""
        self.json_ld_payloads: list[str] = []
        self.spa_payloads: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in self._VOID_TAGS:
            if tag in self._BLOCK_TAGS:
                self._append_boundary()
            return
        attributes = {key.lower(): value or "" for key, value in attrs}
        combined_hint = " ".join(
            (attributes.get("id", ""), attributes.get("class", ""), attributes.get("role", ""))
        )
        hidden = tag in self._HIDDEN_TAGS
        suppressed = tag in self._SUPPRESSED_TAGS or bool(self._SUPPRESSED_HINT.search(combined_hint))
        micro_job = "schema.org/jobposting" in attributes.get("itemtype", "").lower()
        rdfa_job = "jobposting" in attributes.get("typeof", "").lower()
        keys: list[str] = []
        if tag == "main":
            keys.append("main")
        if tag == "article":
            keys.append("article")
        if self._CONTAINER_HINT.search(combined_hint):
            keys.append("container")
        if attributes.get("itemprop", "").lower() in {"description", "jobdescription"} and (
            self._micro_job_depth or micro_job
        ):
            keys.append("microdata")
        if attributes.get("property", "").lower() in {"description", "schema:description"} and (
            self._rdfa_job_depth or rdfa_job
        ):
            keys.append("rdfa")
        frame = {
            "tag": tag,
            "hidden": hidden,
            "suppressed": suppressed,
            "micro_job": micro_job,
            "rdfa_job": rdfa_job,
            "keys": keys,
        }
        self._stack.append(frame)
        self._hidden_depth += int(hidden)
        self._suppressed_depth += int(suppressed)
        self._micro_job_depth += int(micro_job)
        self._rdfa_job_depth += int(rdfa_job)
        for key in keys:
            self._active[key] += 1
        if tag == "script":
            script_type = attributes.get("type", "").lower()
            script_id = attributes.get("id", "").lower()
            if script_type in {"application/ld+json", "application/json"} or script_id in {
                "__next_data__", "__nuxt_data__", "initial_state", "preloaded_state"
            }:
                self._json_script = []
                self._json_script_kind = "json-ld" if script_type == "application/ld+json" else "spa"
        if tag in self._BLOCK_TAGS:
            self._append_boundary()

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in self._BLOCK_TAGS:
            self._append_boundary()

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self._BLOCK_TAGS:
            self._append_boundary()
        if tag == "script" and self._json_script is not None:
            payload = "".join(self._json_script).strip()
            if payload:
                target = self.json_ld_payloads if self._json_script_kind == "json-ld" else self.spa_payloads
                target.append(payload)
            self._json_script = None
            self._json_script_kind = ""
        if not self._stack:
            return
        frame = self._stack.pop()
        self._hidden_depth = max(0, self._hidden_depth - int(bool(frame["hidden"])))
        self._suppressed_depth = max(0, self._suppressed_depth - int(bool(frame["suppressed"])))
        self._micro_job_depth = max(0, self._micro_job_depth - int(bool(frame["micro_job"])))
        self._rdfa_job_depth = max(0, self._rdfa_job_depth - int(bool(frame["rdfa_job"])))
        for key in frame["keys"]:
            self._active[str(key)] = max(0, self._active[str(key)] - 1)

    def handle_data(self, data: str) -> None:
        if self._json_script is not None:
            self._json_script.append(data)
        if self._hidden_depth:
            return
        self.visible.append(data)
        if self._suppressed_depth:
            return
        for key, depth in self._active.items():
            if depth:
                self.candidates[key].append(data)

    def _append_boundary(self) -> None:
        if not self._hidden_depth:
            self.visible.append("\n")
        if not self._hidden_depth and not self._suppressed_depth:
            for key, depth in self._active.items():
                if depth:
                    self.candidates[key].append("\n")


def _html_fragment_text(value: str) -> str:
    parser = _VisibleTextParser()
    parser.feed(value)
    return parser.text()


def _job_posting_descriptions(value: object) -> list[str]:
    results: list[str] = []
    if isinstance(value, list):
        for item in value:
            results.extend(_job_posting_descriptions(item))
    elif isinstance(value, dict):
        raw_type = value.get("@type")
        types = raw_type if isinstance(raw_type, list) else [raw_type]
        if any(str(item).lower() == "jobposting" for item in types):
            description = value.get("description")
            if isinstance(description, str):
                results.append(_html_fragment_text(description))
        for item in value.values():
            results.extend(_job_posting_descriptions(item))
    return results


def _spa_descriptions(value: object) -> list[str]:
    results: list[str] = []
    if isinstance(value, list):
        for item in value:
            results.extend(_spa_descriptions(item))
    elif isinstance(value, dict):
        for key, item in value.items():
            normalized_key = re.sub(r"[^a-z]", "", str(key).lower())
            if normalized_key in {"jobdescription", "vacancydescription"} and isinstance(item, str):
                results.append(_html_fragment_text(item))
            else:
                results.extend(_spa_descriptions(item))
    return results


def extract_vacancy_html(content: bytes) -> VacancyExtraction:
    parser = _VacancyHTMLParser()
    try:
        parser.feed(content.decode("utf-8-sig", errors="replace"))
    except Exception as exc:
        raise ImportError("Could not extract text from HTML document") from exc
    raw_text = normalize_text(" ".join(parser.visible))
    structured: list[tuple[str, list[str]]] = []
    for payload in parser.json_ld_payloads:
        try:
            structured.extend(("json-ld", [item]) for item in _job_posting_descriptions(json.loads(payload)))
        except (json.JSONDecodeError, TypeError, ValueError):
            continue
    if structured:
        method, values = max(structured, key=lambda item: len(normalize_text("\n".join(item[1]))))
        candidate = normalize_text("\n".join(values))
        if candidate:
            return VacancyExtraction(raw_text or candidate, candidate, "html", method)
    for method in ("microdata", "rdfa"):
        candidate = normalize_text(" ".join(parser.candidates[method]))
        if candidate:
            return VacancyExtraction(raw_text or candidate, candidate, "html", method)
    spa_values: list[str] = []
    for payload in parser.spa_payloads:
        try:
            spa_values.extend(_spa_descriptions(json.loads(payload)))
        except (json.JSONDecodeError, TypeError, ValueError):
            continue
    if spa_values:
        candidate = max((normalize_text(item) for item in spa_values), key=len, default="")
        if candidate:
            return VacancyExtraction(raw_text or candidate, candidate, "html", "spa-state")
    for method in ("main", "article"):
        candidate = normalize_text(" ".join(parser.candidates[method]))
        if candidate:
            return VacancyExtraction(raw_text or candidate, candidate, "html", method)
    container_candidate = normalize_text(" ".join(parser.candidates["container"]))
    candidate = container_candidate or raw_text
    if not candidate:
        raise ImportError("No readable text was found in the document")
    method = "probable-container" if container_candidate else "visible-text"
    return VacancyExtraction(raw_text or candidate, candidate, "html", method)


def normalize_text(value: str) -> str:
    lines = []
    for raw_line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t]+", " ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines).strip()


def safe_filename(filename: str) -> str:
    basename = Path(filename).name
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "_", basename).strip(" .")
    return cleaned[:180] or "upload"


def extract_text(filename: str, content: bytes) -> tuple[str, str]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportError("File is larger than the 25 MB MVP limit")
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ImportError(f"Unsupported file type. Supported: {supported}")
    try:
        if extension in {".txt", ".md"}:
            text = content.decode("utf-8-sig")
        elif extension in {".html", ".htm"}:
            parser = _VisibleTextParser()
            parser.feed(content.decode("utf-8-sig", errors="replace"))
            text = parser.text()
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
    except Exception as exc:
        raise ImportError(f"Could not extract text from {extension} document") from exc
    normalized = normalize_text(text)
    if not normalized:
        raise ImportError("No readable text was found in the document")
    return normalized, extension.lstrip(".")


def extract_vacancy_text(filename: str, content: bytes) -> VacancyExtraction:
    extension = Path(filename).suffix.lower()
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportError("File is larger than the 25 MB MVP limit")
    if extension in {".html", ".htm"}:
        return extract_vacancy_html(content)
    text, source_type = extract_text(filename, content)
    return VacancyExtraction(text, text, source_type, f"{source_type}-text")


def write_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
