from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from careerflow.api import create_app
from careerflow.config import Settings
from careerflow.files import extract_text, safe_filename


class VacancyApiTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        settings = Settings(
            root, root / "data", root / "storage", root / "data" / "test.db", "codex", 10
        )
        self.context = TestClient(create_app(settings))
        self.client = self.context.__enter__()
        self.settings = settings

    def tearDown(self) -> None:
        self.context.__exit__(None, None, None)
        self.temp_dir.cleanup()

    def test_create_edit_archive_and_restore_text_vacancy(self) -> None:
        created = self.client.post(
            "/api/vacancies",
            data={
                "text": "Senior Python Engineer\nCompany: Acme Labs\nSalary: €90,000\nBuild reliable APIs.",
                "source": "Referral",
            },
        )
        self.assertEqual(created.status_code, 201)
        vacancy = created.json()
        self.assertEqual(vacancy["title"], "Senior Python Engineer")
        self.assertEqual(vacancy["company"], "Acme Labs")
        self.assertIn("90,000", vacancy["salary"])
        self.assertTrue(Path(vacancy["extracted_text_path"]).exists())

        updated = self.client.patch(
            f"/api/vacancies/{vacancy['id']}", json={"company": "Acme International"}
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()["company"], "Acme International")

        self.assertEqual(
            self.client.post(f"/api/vacancies/{vacancy['id']}/archive").status_code, 200
        )
        self.assertEqual(self.client.get("/api/vacancies").json(), [])
        self.assertEqual(len(self.client.get("/api/vacancies?archived=true").json()), 1)
        self.assertEqual(
            self.client.post(f"/api/vacancies/{vacancy['id']}/restore").status_code, 200
        )
        self.assertEqual(len(self.client.get("/api/vacancies").json()), 1)

    def test_description_patch_rejects_empty_text_without_mutation(self) -> None:
        vacancy = self.client.post(
            "/api/vacancies", data={"text": "Engineer\nPython required"}
        ).json()
        url = f"/api/vacancies/{vacancy['id']}"
        rejected = self.client.patch(url, json={"description": "  \n  \t "})
        self.assertEqual(rejected.status_code, 422)
        self.assertEqual(self.client.get(url).json()["description"], vacancy["description"])

        updated = self.client.patch(url, json={"description": "  Required Python\n\n  SQL  "})
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()["description"], "Required Python\nSQL")

    def test_html_import_removes_scripts_and_preserves_original(self) -> None:
        content = b"<html><style>.bad{}</style><h1>Backend Engineer</h1><p>Company: Northwind</p><script>alert(1)</script><p>Python and SQL</p></html>"
        response = self.client.post(
            "/api/vacancies", files={"file": ("job.html", content, "text/html")}
        )
        self.assertEqual(response.status_code, 201)
        vacancy = response.json()
        self.assertIn("Backend Engineer", vacancy["description"])
        self.assertNotIn("alert", vacancy["description"])
        self.assertTrue(Path(vacancy["original_path"]).exists())

    def test_docx_extraction_and_invalid_extension(self) -> None:
        path = Path(self.temp_dir.name) / "job.docx"
        document = Document()
        document.add_heading("Product Manager", level=1)
        document.add_paragraph("Company: Contoso")
        document.save(path)
        text, source_type = extract_text(path.name, path.read_bytes())
        self.assertEqual(source_type, "docx")
        self.assertIn("Product Manager", text)

        invalid = self.client.post(
            "/api/vacancies", files={"file": ("job.exe", b"data", "application/octet-stream")}
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("Unsupported", invalid.json()["detail"])

    def test_missing_vacancy_and_filename_safety(self) -> None:
        self.assertEqual(self.client.get("/api/vacancies/999").status_code, 404)
        self.assertEqual(safe_filename("../../unsafe?.txt"), "unsafe_.txt")


if __name__ == "__main__":
    unittest.main()
